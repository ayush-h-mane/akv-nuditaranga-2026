import io
import csv
import datetime
from typing import List, Optional
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query, Response, status
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from sqlalchemy import func, or_
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from ..database import get_db
from ..models import User, Admin, Event, Registration, VolunteerAttendance, AuditLog
from ..auth_deps import require_superadmin
from ..services.email_service import send_admin_approval_email

router = APIRouter(prefix="/superadmin", tags=["Super Admin"])

# Schemas
class StudentUpdateRequest(BaseModel):
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    institute: Optional[str] = None
    department: Optional[str] = None
    role: Optional[str] = None  # VOLUNTEER, PARTICIPANT, SPECTATOR
    account_status: Optional[str] = None  # ACTIVE, DISABLED

class AttendanceUpdateRequest(BaseModel):
    status: str  # PRESENT, ABSENT, LATE, EXCUSED
    notes: Optional[str] = None

class AttendanceCreateRequest(BaseModel):
    volunteer_user_id: int
    date: str  # YYYY-MM-DD
    status: str = "PRESENT"
    notes: Optional[str] = None

# ==========================================
# 1. OVERVIEW & METRICS
# ==========================================
@router.get("/stats")
def get_superadmin_stats(
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    total_students = db.query(func.count(User.id)).filter(
        User.role.in_(["STUDENT", "VOLUNTEER", "PARTICIPANT", "SPECTATOR"])
    ).scalar() or 0

    total_volunteers = db.query(func.count(User.id)).filter(User.role == "VOLUNTEER").scalar() or 0
    total_participants = db.query(func.count(User.id)).filter(User.role == "PARTICIPANT").scalar() or 0
    total_spectators = db.query(func.count(User.id)).filter(User.role == "SPECTATOR").scalar() or 0
    total_events = db.query(func.count(Event.id)).scalar() or 0
    total_registrations = db.query(func.count(Registration.id)).scalar() or 0

    pending_admins = db.query(func.count(Admin.id)).filter(
        Admin.approval_status == "PENDING_APPROVAL"
    ).scalar() or 0

    today_str = datetime.date.today().strftime("%Y-%m-%d")
    today_present = db.query(func.count(VolunteerAttendance.id)).filter(
        VolunteerAttendance.date == today_str,
        VolunteerAttendance.status == "PRESENT"
    ).scalar() or 0

    today_absent = db.query(func.count(VolunteerAttendance.id)).filter(
        VolunteerAttendance.date == today_str,
        VolunteerAttendance.status == "ABSENT"
    ).scalar() or 0

    return {
        "success": True,
        "metrics": {
            "total_students": total_students,
            "total_volunteers": total_volunteers,
            "total_participants": total_participants,
            "total_spectators": total_spectators,
            "total_events": total_events,
            "total_registrations": total_registrations,
            "pending_admins": pending_admins,
            "today_attendance": {
                "date": today_str,
                "present": today_present,
                "absent": today_absent,
                "total_marked": today_present + today_absent,
                "total_volunteers": total_volunteers
            }
        }
    }

# ==========================================
# 2. STUDENT MANAGEMENT (CRUD)
# ==========================================
@router.get("/students")
def list_students(
    search: Optional[str] = None,
    role: Optional[str] = None,
    status_filter: Optional[str] = None,
    department: Optional[str] = None,
    limit: int = 100,
    offset: int = 0,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    query = db.query(User).filter(User.role != "SUPERADMIN")

    if role and role.upper() != "ALL":
        query = query.filter(User.role == role.upper())
    if status_filter and status_filter.upper() != "ALL":
        query = query.filter(User.account_status == status_filter.upper())
    if department and department != "all":
        query = query.filter(User.department == department)
    if search:
        s = f"%{search.strip().lower()}%"
        query = query.filter(
            or_(
                func.lower(User.name).like(s),
                func.lower(User.auid).like(s),
                func.lower(User.email).like(s),
                func.lower(User.phone).like(s),
                func.lower(User.department).like(s),
                func.lower(User.registration_id).like(s)
            )
        )

    total_count = query.count()
    users = query.order_by(User.created_at.desc()).offset(offset).limit(limit).all()

    student_list = []
    for u in users:
        reg_count = db.query(func.count(Registration.id)).filter(
            or_(Registration.user_id == u.id, Registration.auid == u.auid)
        ).scalar() or 0

        student_list.append({
            "id": u.id,
            "name": u.name,
            "auid": u.auid,
            "email": u.email,
            "phone": u.phone,
            "institute": u.institute,
            "department": u.department,
            "semester": u.semester,
            "section": u.section,
            "gender": u.gender,
            "role": u.role,
            "registration_id": u.registration_id,
            "account_status": u.account_status,
            "event_registrations_count": reg_count,
            "created_at": u.created_at.isoformat() if u.created_at else None
        })

    return {
        "total": total_count,
        "students": student_list
    }

@router.put("/students/{user_id}")
def update_student(
    user_id: int,
    payload: StudentUpdateRequest,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    student = db.query(User).filter(User.id == user_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    old_details = f"Role: {student.role}, Status: {student.account_status}, Name: {student.name}"
    
    if payload.name:
        student.name = payload.name.strip()
    if payload.email:
        student.email = payload.email.strip().lower()
    if payload.phone:
        student.phone = payload.phone.strip()
    if payload.institute:
        student.institute = payload.institute.strip()
    if payload.department:
        student.department = payload.department.strip()
    if payload.role and payload.role.upper() in ["VOLUNTEER", "PARTICIPANT", "SPECTATOR", "STUDENT"]:
        student.role = payload.role.upper()
    if payload.account_status and payload.account_status.upper() in ["ACTIVE", "DISABLED"]:
        student.account_status = payload.account_status.upper()

    student.updated_at = datetime.datetime.utcnow()
    new_details = f"Role: {student.role}, Status: {student.account_status}, Name: {student.name}"

    log = AuditLog(
        user_id=current_user.id,
        actor_name=current_user.name,
        action="STUDENT_UPDATED",
        target_type="STUDENT",
        target_id=str(student.id),
        previous_value=old_details,
        new_value=new_details
    )
    db.add(log)
    db.commit()

    return {"success": True, "message": "Student information updated successfully"}

@router.delete("/students/{user_id}")
def delete_student(
    user_id: int,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    student = db.query(User).filter(User.id == user_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    if student.role == "SUPERADMIN":
        raise HTTPException(status_code=403, detail="Cannot delete Super Admin account")

    student_name = student.name
    student_auid = student.auid

    # Log audit before deletion
    log = AuditLog(
        user_id=current_user.id,
        actor_name=current_user.name,
        action="STUDENT_DELETED",
        target_type="STUDENT",
        target_id=str(user_id),
        previous_value=f"Name: {student_name}, AUID: {student_auid}",
        new_value="DELETED"
    )
    db.add(log)

    db.delete(student)
    db.commit()

    return {"success": True, "message": f"Student '{student_name}' deleted successfully"}

# ==========================================
# 3. ADMIN APPROVAL & MANAGEMENT
# ==========================================
@router.get("/admins")
def list_admins(
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    # Retrieve all registered admins, surfacing PENDING_APPROVAL at the very top
    admins = db.query(Admin).order_by(
        (Admin.approval_status == "PENDING_APPROVAL").desc(),
        Admin.created_at.desc()
    ).all()

    # Safety check: Detect any User marked role="ADMIN" without an Admin profile and backfill
    existing_user_ids = {a.user_id for a in admins if a.user_id}
    orphan_admins = db.query(User).filter(
        User.role == "ADMIN",
        ~User.id.in_(existing_user_ids) if existing_user_ids else True
    ).all()

    for ou in orphan_admins:
        try:
            uname = ou.auid.lower().replace("-", "_") if ou.auid else f"admin_{ou.id}"
            if db.query(Admin).filter(func.lower(Admin.username) == uname.lower()).first():
                uname = f"{uname}_{ou.id}"
            new_a = Admin(
                user_id=ou.id,
                username=uname,
                admin_type=ou.admin_type or "WORKING_COMMITTEE",
                faculty_id=ou.faculty_id,
                approval_status="PENDING_APPROVAL",
                created_at=ou.created_at or datetime.datetime.utcnow()
            )
            db.add(new_a)
            db.commit()
            db.refresh(new_a)
            admins.insert(0, new_a)
        except Exception as e:
            db.rollback()
            print(f"[ORPHAN ADMIN BACKFILL WARNING] {e}")

    results = []
    for a in admins:
        u = a.user
        results.append({
            "id": a.id,
            "user_id": a.user_id,
            "username": a.username,
            "full_name": u.name if u else "N/A",
            "email": u.email if u else "N/A",
            "phone": u.phone if u else "N/A",
            "institute": u.institute if u else "N/A",
            "department": u.department if u else "N/A",
            "admin_type": a.admin_type or (u.admin_type if u else "WORKING_COMMITTEE"),
            "faculty_id": a.faculty_id or (u.faculty_id if u else None),
            "photo_url": u.photo_url if u else None,
            "role": u.role if u else "ADMIN",
            "approval_status": a.approval_status,
            "approved_by": a.approved_by,
            "approved_at": a.approved_at.isoformat() if a.approved_at else None,
            "account_status": u.account_status if u else "N/A",
            "created_at": a.created_at.isoformat() if a.created_at else None
        })
    return results

@router.post("/admins/{admin_id}/approve")
def approve_admin(
    admin_id: int,
    background_tasks: BackgroundTasks,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    admin = db.query(Admin).filter(Admin.id == admin_id).first()
    if not admin:
        raise HTTPException(status_code=404, detail="Admin record not found")

    old_status = admin.approval_status
    admin.approval_status = "APPROVED"
    admin.approved_by = current_user.name
    admin.approved_at = datetime.datetime.utcnow()

    # Ensure user account is active
    if admin.user:
        admin.user.account_status = "ACTIVE"
        if admin.user.email:
            background_tasks.add_task(send_admin_approval_email, admin.user.email, admin.user.name, "APPROVED")

    log = AuditLog(
        user_id=current_user.id,
        actor_name=current_user.name,
        action="ADMIN_APPROVED",
        target_type="ADMIN",
        target_id=str(admin.id),
        previous_value=old_status,
        new_value="APPROVED"
    )
    db.add(log)
    db.commit()

    return {"success": True, "message": f"Admin '{admin.username}' has been approved."}

@router.post("/admins/{admin_id}/reject")
def reject_admin(
    admin_id: int,
    background_tasks: BackgroundTasks,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    admin = db.query(Admin).filter(Admin.id == admin_id).first()
    if not admin:
        raise HTTPException(status_code=404, detail="Admin record not found")

    old_status = admin.approval_status
    admin.approval_status = "REJECTED"
    admin.approved_by = current_user.name

    if admin.user and admin.user.email:
        background_tasks.add_task(send_admin_approval_email, admin.user.email, admin.user.name, "REJECTED")

    log = AuditLog(
        user_id=current_user.id,
        actor_name=current_user.name,
        action="ADMIN_REJECTED",
        target_type="ADMIN",
        target_id=str(admin.id),
        previous_value=old_status,
        new_value="REJECTED"
    )
    db.add(log)
    db.commit()

    return {"success": True, "message": f"Admin '{admin.username}' has been rejected."}

@router.post("/admins/{admin_id}/toggle-status")
def toggle_admin_status(
    admin_id: int,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    admin = db.query(Admin).filter(Admin.id == admin_id).first()
    if not admin or not admin.user:
        raise HTTPException(status_code=404, detail="Admin record not found")

    new_status = "DISABLED" if admin.user.account_status == "ACTIVE" else "ACTIVE"
    old_status = admin.user.account_status
    admin.user.account_status = new_status

    log = AuditLog(
        user_id=current_user.id,
        actor_name=current_user.name,
        action="ADMIN_STATUS_CHANGED",
        target_type="ADMIN",
        target_id=str(admin.id),
        previous_value=f"Account status: {old_status}",
        new_value=f"Account status: {new_status}"
    )
    db.add(log)
    db.commit()

    return {"success": True, "message": f"Admin status updated to {new_status}"}

@router.delete("/admins/{admin_id}")
def delete_admin(
    admin_id: int,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    admin = db.query(Admin).filter(Admin.id == admin_id).first()
    if not admin:
        raise HTTPException(status_code=404, detail="Admin not found")

    username = admin.username
    user = admin.user

    # Requirement 9: Block deleting SuperAdmin account
    if (user and user.role == "SUPERADMIN") or username in ["superadmin", "akv-nt-2026"]:
        raise HTTPException(status_code=403, detail="SuperAdmin profile cannot be deleted.")

    log = AuditLog(
        user_id=current_user.id,
        actor_name=current_user.name,
        action="ADMIN_DELETED",
        target_type="ADMIN",
        target_id=str(admin_id),
        previous_value=f"Admin Username: {username}",
        new_value="DELETED"
    )
    db.add(log)

    db.delete(admin)
    if user and user.role == "ADMIN":
        db.delete(user)
    db.commit()

    return {"success": True, "message": f"Admin '{username}' removed successfully"}

# ==========================================
# 4. VOLUNTEER MANAGEMENT (AUTO-GENERATED)
# ==========================================
@router.get("/volunteers")
def list_all_volunteers(
    department: Optional[str] = None,
    search: Optional[str] = None,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    """
    Volunteers are automatically queried from student registrations where role == 'VOLUNTEER'.
    No manual list creation is allowed or needed.
    """
    query = db.query(User).filter(User.role == "VOLUNTEER")

    if department and department != "all":
        query = query.filter(User.department == department)
    if search:
        s = f"%{search.strip().lower()}%"
        query = query.filter(
            or_(
                func.lower(User.name).like(s),
                func.lower(User.auid).like(s),
                func.lower(User.department).like(s),
                func.lower(User.phone).like(s)
            )
        )

    volunteers = query.order_by(User.name.asc()).all()
    today_str = datetime.date.today().strftime("%Y-%m-%d")

    results = []
    for v in volunteers:
        # Check today's attendance
        today_att = db.query(VolunteerAttendance).filter(
            VolunteerAttendance.user_id == v.id,
            VolunteerAttendance.date == today_str
        ).first()

        # Count total days present
        total_present = db.query(func.count(VolunteerAttendance.id)).filter(
            VolunteerAttendance.user_id == v.id,
            VolunteerAttendance.status == "PRESENT"
        ).scalar() or 0

        results.append({
            "user_id": v.id,
            "name": v.name,
            "auid": v.auid,
            "department": v.department,
            "institute": v.institute,
            "phone": v.phone,
            "email": v.email,
            "semester": v.semester,
            "section": v.section,
            "gender": v.gender,
            "registration_id": v.registration_id,
            "today_attendance": today_att.status if today_att else "NOT_MARKED",
            "today_checkin_time": today_att.check_in_time.strftime("%I:%M %p") if today_att and today_att.check_in_time else None,
            "total_days_present": total_present,
            "registered_at": v.created_at.strftime("%Y-%m-%d") if v.created_at else None
        })

    return results

# ==========================================
# 5. SUPER ADMIN ATTENDANCE (FULL OVERSIGHT)
# ==========================================
@router.get("/attendance")
def get_superadmin_attendance(
    date: Optional[str] = None,  # Specific date YYYY-MM-DD or all
    department: Optional[str] = None,
    search: Optional[str] = None,
    status_filter: Optional[str] = None,
    limit: int = 150,
    offset: int = 0,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    query = db.query(VolunteerAttendance)

    if date and date != "all":
        query = query.filter(VolunteerAttendance.date == date)
    if department and department != "all":
        query = query.filter(VolunteerAttendance.department == department)
    if status_filter and status_filter != "all":
        query = query.filter(VolunteerAttendance.status == status_filter.upper())
    if search:
        s = f"%{search.strip().lower()}%"
        query = query.filter(
            or_(
                func.lower(VolunteerAttendance.volunteer_name).like(s),
                func.lower(VolunteerAttendance.auid).like(s),
                func.lower(VolunteerAttendance.department).like(s)
            )
        )

    records = query.order_by(VolunteerAttendance.date.desc(), VolunteerAttendance.volunteer_name.asc()).offset(offset).limit(limit).all()
    
    # Available unique dates for filter dropdown
    unique_dates = [r[0] for r in db.query(VolunteerAttendance.date).distinct().order_by(VolunteerAttendance.date.desc()).all()]

    output = []
    for r in records:
        output.append({
            "id": r.id,
            "user_id": r.user_id,
            "auid": r.auid,
            "volunteer_name": r.volunteer_name,
            "department": r.department,
            "date": r.date,
            "status": r.status,
            "check_in_time": r.check_in_time.strftime("%I:%M %p") if r.check_in_time else "N/A",
            "marked_by": r.marked_by,
            "last_modified": r.updated_at.strftime("%Y-%m-%d %I:%M %p") if r.updated_at else ""
        })

    return {
        "records": output,
        "available_dates": unique_dates
    }

@router.put("/attendance/{attendance_id}")
def update_attendance_record(
    attendance_id: int,
    payload: AttendanceUpdateRequest,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    rec = db.query(VolunteerAttendance).filter(VolunteerAttendance.id == attendance_id).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Attendance record not found")

    old_status = rec.status
    new_status = payload.status.upper()
    if new_status not in ["PRESENT", "ABSENT", "LATE", "EXCUSED"]:
        raise HTTPException(status_code=400, detail="Invalid attendance status")

    rec.status = new_status
    if new_status == "PRESENT" and not rec.check_in_time:
        rec.check_in_time = datetime.datetime.utcnow()
    elif new_status == "ABSENT":
        rec.check_in_time = None

    rec.marked_by = f"{current_user.name} (Super Admin)"
    rec.updated_at = datetime.datetime.utcnow()

    # Mandatory Audit Log Entry
    log = AuditLog(
        user_id=current_user.id,
        actor_name=current_user.name,
        action="ATTENDANCE_MODIFIED",
        target_type="ATTENDANCE",
        target_id=str(rec.id),
        previous_value=f"Volunteer: {rec.volunteer_name}, Date: {rec.date}, Status: {old_status}",
        new_value=f"Status: {new_status}, Reason/Note: {payload.notes or 'Manual correction'}"
    )
    db.add(log)
    db.commit()

    return {
        "success": True,
        "message": f"Attendance for '{rec.volunteer_name}' on {rec.date} updated from {old_status} to {new_status}."
    }

@router.post("/attendance/mark")
def create_or_mark_attendance(
    payload: AttendanceCreateRequest,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    volunteer = db.query(User).filter(User.id == payload.volunteer_user_id).first()
    if not volunteer:
        raise HTTPException(status_code=404, detail="Volunteer not found")

    status_val = payload.status.upper()
    if status_val not in ["PRESENT", "ABSENT", "LATE", "EXCUSED"]:
        raise HTTPException(status_code=400, detail="Invalid status")

    existing = db.query(VolunteerAttendance).filter(
        VolunteerAttendance.user_id == volunteer.id,
        VolunteerAttendance.date == payload.date
    ).first()

    if existing:
        old_val = existing.status
        existing.status = status_val
        if status_val == "PRESENT" and not existing.check_in_time:
            existing.check_in_time = datetime.datetime.utcnow()
        elif status_val == "ABSENT":
            existing.check_in_time = None
        existing.marked_by = f"{current_user.name} (Super Admin)"
        existing.updated_at = datetime.datetime.utcnow()

        log = AuditLog(
            user_id=current_user.id,
            actor_name=current_user.name,
            action="ATTENDANCE_OVERWRITE",
            target_type="ATTENDANCE",
            target_id=str(existing.id),
            previous_value=f"Status: {old_val}",
            new_value=f"Status: {status_val}"
        )
        db.add(log)
    else:
        new_rec = VolunteerAttendance(
            user_id=volunteer.id,
            auid=volunteer.auid,
            volunteer_name=volunteer.name,
            department=volunteer.department,
            date=payload.date,
            status=status_val,
            check_in_time=datetime.datetime.utcnow() if status_val == "PRESENT" else None,
            marked_by=f"{current_user.name} (Super Admin)"
        )
        db.add(new_rec)
        db.flush()

        log = AuditLog(
            user_id=current_user.id,
            actor_name=current_user.name,
            action="ATTENDANCE_CREATED",
            target_type="ATTENDANCE",
            target_id=str(new_rec.id),
            previous_value="None",
            new_value=f"Volunteer: {volunteer.name}, Date: {payload.date}, Status: {status_val}"
        )
        db.add(log)

    db.commit()
    return {"success": True, "message": f"Attendance marked for {volunteer.name} ({payload.date}) as {status_val}"}

# ==========================================
# 6. ATTENDANCE EXPORTS (CSV & XLSX)
# ==========================================
@router.get("/attendance/export-csv")
def export_attendance_csv(
    date: Optional[str] = None,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    query = db.query(VolunteerAttendance)
    if date and date != "all":
        query = query.filter(VolunteerAttendance.date == date)
    
    records = query.order_by(VolunteerAttendance.date.desc(), VolunteerAttendance.volunteer_name.asc()).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Record ID",
        "AUID",
        "Volunteer Name",
        "Department",
        "Date",
        "Attendance Status",
        "Check-In Time",
        "Marked By",
        "Last Modified"
    ])

    for r in records:
        writer.writerow([
            r.id,
            r.auid,
            r.volunteer_name,
            r.department,
            r.date,
            r.status,
            r.check_in_time.strftime("%I:%M %p") if r.check_in_time else "N/A",
            r.marked_by,
            r.updated_at.strftime("%Y-%m-%d %H:%M:%S") if r.updated_at else ""
        ])

    csv_data = output.getvalue()
    filename = f"AKV_Nuditaranga_2026_Volunteer_Attendance_{date or 'All'}.csv"
    return Response(
        content=csv_data,
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/attendance/export-xlsx")
def export_attendance_xlsx(
    date: Optional[str] = None,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    query = db.query(VolunteerAttendance)
    if date and date != "all":
        query = query.filter(VolunteerAttendance.date == date)
    
    records = query.order_by(VolunteerAttendance.date.desc(), VolunteerAttendance.volunteer_name.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Volunteer Attendance"

    # Styling Palettes
    header_fill = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")  # Karnataka Red
    sub_fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")     # Yellow accent
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=15, bold=True, color="991B1B")
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Title block
    ws.merge_cells("A1:H1")
    ws["A1"] = "ACHARYA KANNADA VEDIKE (AKV) — NUDITARANGA 2026"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:H2")
    date_label = f"Filter Date: {date}" if date and date != "all" else "Complete Attendance Record"
    ws["A2"] = f"Official Volunteer Daily Attendance Report • {date_label} • Generated by: {current_user.name}"
    ws["A2"].font = Font(size=10, italic=True, color="57534E")
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.append([])  # Blank row

    # Headers
    headers = [
        "Record ID",
        "AUID",
        "Volunteer Full Name",
        "Department",
        "Date",
        "Status",
        "Check-In Time",
        "Marked By"
    ]
    ws.append(headers)
    header_row = 4

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = header_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for r in records:
        checkin_val = r.check_in_time.strftime("%I:%M %p") if r.check_in_time else "N/A"
        row_data = [
            r.id,
            r.auid,
            r.volunteer_name,
            r.department,
            r.date,
            r.status,
            checkin_val,
            r.marked_by
        ]
        ws.append(row_data)

    # Format data rows
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")
            if cell.column == 6:  # Status column
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if cell.value == "PRESENT":
                    cell.font = Font(bold=True, color="15803D")
                elif cell.value == "ABSENT":
                    cell.font = Font(bold=True, color="B91C1C")

    # Adjust column widths
    column_widths = {
        "A": 12,
        "B": 18,
        "C": 28,
        "D": 30,
        "E": 14,
        "F": 15,
        "G": 18,
        "H": 25
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Save to stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "AKV_Nuditaranga_2026_Volunteer_Attendance.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# ==========================================
# 7. AUDIT LOGS
# ==========================================
@router.get("/audit-logs")
def get_audit_logs(
    limit: int = 100,
    offset: int = 0,
    action: Optional[str] = None,
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    query = db.query(AuditLog)
    if action and action != "all":
        query = query.filter(AuditLog.action == action)
    
    logs = query.order_by(AuditLog.created_at.desc()).offset(offset).limit(limit).all()
    
    results = []
    for l in logs:
        results.append({
            "id": l.id,
            "actor_name": l.actor_name,
            "action": l.action,
            "target_type": l.target_type,
            "target_id": l.target_id,
            "previous_value": l.previous_value,
            "new_value": l.new_value,
            "timestamp": l.created_at.strftime("%Y-%m-%d %I:%M:%S %p") if l.created_at else ""
        })
    return results

# ==========================================
# 8. DATABASE ERASE & RESET ENDPOINT
# ==========================================
@router.post("/database/reset")
def reset_database_endpoint(
    current_user: User = Depends(require_superadmin),
    db: Session = Depends(get_db)
):
    """
    Super Admin endpoint to safely reset the database and re-seed defaults.
    """
    from ...reset_db import erase_and_reset_database
    try:
        erase_and_reset_database()
        return {
            "success": True,
            "message": "Database has been completely erased, reseeded, and vacuumed successfully."
        }
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to reset database: {str(e)}"
        )

